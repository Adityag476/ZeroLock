import os
import sys
import io
import time
import json

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')
import requests
import sqlite3
from docx import Document
from PIL import Image
import fitz

BASE_URL = "http://localhost:8000"
results = {}

print("=== Starting Section B Audit ===")

def get_exam_count():
    conn = sqlite3.connect("zeroleak.db")
    c = conn.cursor()
    c.execute("SELECT count(*) FROM exams")
    cnt = c.fetchone()[0]
    conn.close()
    return cnt

future_release = time.time() + 3600.0

# B1: Upload TY_DSP_26-27_Tutorial_6.pdf
print("\n--- B1: Tutorial 6 upload ---")
t6_path = "experiments/fixtures/TY_DSP_26-27_Tutorial_6.pdf"
with open(t6_path, "rb") as f:
    r = requests.post(
        f"{BASE_URL}/api/exams/",
        files={"file": ("TY_DSP_26-27_Tutorial_6.pdf", f, "application/pdf")},
        data={"name": "DSP Tutorial 6 Audit", "release_time": str(future_release)}
    )
print(f"Status: {r.status_code}")
b1_data = r.json() if r.status_code == 200 else r.text
print(f"Response: {b1_data}")
b1_pass = False
if r.status_code == 200:
    q_count = b1_data.get("question_count", 0)
    engine = b1_data.get("engine", "")
    preview = b1_data.get("preview", "")
    has_text = "impulse-invariant" in preview.lower()
    b1_pass = (q_count == 4) and (engine == "compact") and has_text
    print(f"q_count={q_count} (exp 4), engine={engine} (exp compact), has_text={has_text}")
results["B1"] = {"pass": b1_pass, "status": r.status_code, "data": b1_data}

# B2: Upload Tutorial 4
print("\n--- B2: Tutorial 4 upload ---")
t4_path = "experiments/fixtures/TY_DSP_26-27_Tutorial_4.pdf"
with open(t4_path, "rb") as f:
    r = requests.post(
        f"{BASE_URL}/api/exams/",
        files={"file": ("TY_DSP_26-27_Tutorial_4.pdf", f, "application/pdf")},
        data={"name": "DSP Tutorial 4 Audit", "release_time": str(future_release)}
    )
print(f"Status: {r.status_code}")
b2_data = r.json() if r.status_code == 200 else r.text
print(f"Response: {b2_data}")
b2_pass = False
if r.status_code == 200:
    q_count = b2_data.get("question_count", 0)
    b2_pass = (q_count == 12)
    print(f"q_count={q_count} (exp 12)")
results["B2"] = {"pass": b2_pass, "status": r.status_code, "data": b2_data}

# B3: Upload one-sentence .txt
print("\n--- B3: One-sentence .txt ---")
count_before = get_exam_count()
txt_bytes = b"This is a single short sentence for an exam."
r = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("short.txt", io.BytesIO(txt_bytes), "text/plain")},
    data={"name": "Short Exam", "release_time": str(future_release)}
)
print(f"Status: {r.status_code}")
print(f"Response text: {r.text}")
count_after = get_exam_count()
b3_pass = (r.status_code == 400) and ("capacity" in r.text.lower() or "too short" in r.text.lower() or "minimum" in r.text.lower()) and (count_before == count_after)
print(f"count_before={count_before}, count_after={count_after}")
results["B3"] = {"pass": b3_pass, "status": r.status_code, "text": r.text, "count_diff": count_after - count_before}

# B4: Upload image-only PDF
print("\n--- B4: Scanned / Image-only PDF ---")
img = Image.new("RGB", (600, 800), color=(240, 240, 240))
pdf_doc = fitz.open()
img_bytes_io = io.BytesIO()
img.save(img_bytes_io, format="PNG")
img_bytes = img_bytes_io.getvalue()
page = pdf_doc.new_page(width=600, height=800)
page.insert_image(page.rect, stream=img_bytes)
scanned_pdf_bytes = pdf_doc.write()
pdf_doc.close()

r = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("scanned_exam.pdf", io.BytesIO(scanned_pdf_bytes), "application/pdf")},
    data={"name": "Scanned Exam", "release_time": str(future_release)}
)
print(f"Status: {r.status_code}")
print(f"Response text: {r.text}")
b4_pass = (r.status_code == 400) and ("text-based" in r.text.lower())
results["B4"] = {"pass": b4_pass, "status": r.status_code, "text": r.text}

# B5: Upload a .png renamed to .pdf
print("\n--- B5: PNG renamed to .pdf ---")
r = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("fake_pdf.pdf", io.BytesIO(img_bytes), "application/pdf")},
    data={"name": "Fake PDF Exam", "release_time": str(future_release)}
)
print(f"Status: {r.status_code}")
print(f"Response text: {r.text}")
b5_pass = (r.status_code == 400)
results["B5"] = {"pass": b5_pass, "status": r.status_code, "text": r.text}

# B6: Upload .docx with 3 questions
print("\n--- B6: DOCX with 3 questions ---")
doc = Document()
doc.add_heading("Midterm Exam", level=1)
doc.add_paragraph("1. Explain the sampling theorem and Nyquist rate in discrete-time signal processing and communication systems.")
doc.add_paragraph("2. Derive the Discrete Fourier Transform for an arbitrary N-point discrete sequence with mathematical justification.")
doc.add_paragraph("3. Compare finite impulse response (FIR) and infinite impulse response (IIR) digital filter design methodologies in detail.")
docx_io = io.BytesIO()
doc.save(docx_io)
docx_bytes = docx_io.getvalue()

r = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("questions.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    data={"name": "DOCX Exam", "release_time": str(future_release)}
)
print(f"Status: {r.status_code}")
b6_data = r.json() if r.status_code == 200 else r.text
print(f"Response: {b6_data}")
b6_pass = False
if r.status_code == 200:
    q_count = b6_data.get("question_count", 0)
    b6_pass = (q_count == 3)
    print(f"q_count={q_count} (exp 3)")
results["B6"] = {"pass": b6_pass, "status": r.status_code, "data": b6_data}

# B7: Re-upload same file twice -> two distinct exams, both functional
print("\n--- B7: Duplicate upload ---")
with open(t6_path, "rb") as f:
    content = f.read()

r1 = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("T6_dup1.pdf", io.BytesIO(content), "application/pdf")},
    data={"name": "DSP T6 Run 1", "release_time": str(future_release)}
)
r2 = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("T6_dup2.pdf", io.BytesIO(content), "application/pdf")},
    data={"name": "DSP T6 Run 2", "release_time": str(future_release)}
)
b7_pass = False
if r1.status_code == 200 and r2.status_code == 200:
    id1 = r1.json().get("exam_id")
    id2 = r2.json().get("exam_id")
    distinct = (id1 != id2 and id1 is not None and id2 is not None)
    g1 = requests.get(f"{BASE_URL}/api/exams/{id1}")
    g2 = requests.get(f"{BASE_URL}/api/exams/{id2}")
    b7_pass = distinct and (g1.status_code == 200) and (g2.status_code == 200)
    print(f"id1={id1}, id2={id2}, distinct={distinct}, g1={g1.status_code}, g2={g2.status_code}")
results["B7"] = {"pass": b7_pass, "r1": r1.status_code, "r2": r2.status_code}

print("\n=== SUMMARY OF SECTION B ===")
for k, v in results.items():
    print(f"{k}: {'PASS' if v['pass'] else 'FAIL'} | {v}")
