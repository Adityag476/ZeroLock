import os
import sys
import io
import re
import time
import json
import random
import requests
import sqlite3
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from core.variant_engine import generate_centre_variant, SWAP_GROUPS

BASE_URL = "http://localhost:8000"
results = {}

print("=== Starting Section E Audit ===")

# Create exam for Section E with numerical and wording variation capabilities
doc = Document()
doc.add_heading("Physics and Electrical Engineering", level=1)
doc.add_paragraph("1. A projectile is launched with velocity 50 m/s. Calculate the maximum height reached by the object.")
doc.add_paragraph("2. An AC circuit has frequency 50 Hz and resistance 100 ohm. Find the total impedance of the circuit.")
doc.add_paragraph("3. A mass of 10 kg moves at speed 20 m/s. Compute the total kinetic energy of the system.")
doc.add_paragraph("4. A sample with volume 2 L is held at 300 K. Show that the pressure is approximately constant under uniform heating.")
docx_io = io.BytesIO()
doc.save(docx_io)

r_ex = requests.post(
    f"{BASE_URL}/api/exams/",
    files={"file": ("physics_exam.docx", docx_io.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    data={"name": "Physics and Electrical Engineering", "release_time": str(time.time() + 3600.0)}
)
exam_id = r_ex.json()["exam_id"]
print(f"Exam created for Section E: {exam_id}")

conn = sqlite3.connect("zeroleak.db")
c = conn.cursor()
c.execute("SELECT questions_json FROM exams WHERE id = ?", (exam_id,))
questions = json.loads(c.fetchone()[0])
conn.close()

answer_key = "Q1: 50 m/s\nQ2: 50 Hz, 100 ohm\nQ3: 10 kg, 20 m/s\nQ4: 2 L, 300 K"

# E1: Variants ON: two centres' texts differ ONLY in sanctioned slots (diff check)
print("\n--- E1: Diff check between centre 14 and centre 28 ---")
v14 = generate_centre_variant(exam_id, 14, questions, answer_key)
v28 = generate_centre_variant(exam_id, 28, questions, answer_key)

e1_clean = True
swapped_words = set()
for g in SWAP_GROUPS:
    for phrase in g:
        for w in phrase.split():
            swapped_words.add(w.lower())

for q1, q2 in zip(v14["variant_questions"], v28["variant_questions"]):
    tokens1 = re.findall(r'\b\w+\b', q1)
    tokens2 = re.findall(r'\b\w+\b', q2)
    # Check words present in q1 but not q2
    for t in tokens1:
        if t not in tokens2:
            is_num = any(ch.isdigit() for ch in t)
            is_swap = t.lower() in swapped_words
            if not (is_num or is_swap):
                print(f"Non-sanctioned diff token in q1: '{t}'")
                e1_clean = False
    for t in tokens2:
        if t not in tokens1:
            is_num = any(ch.isdigit() for ch in t)
            is_swap = t.lower() in swapped_words
            if not (is_num or is_swap):
                print(f"Non-sanctioned diff token in q2: '{t}'")
                e1_clean = False

e1_pass = e1_clean and len(v14["swap_vector"]) > 0 and len(v28["swap_vector"]) > 0
print(f"E1 diff check clean: {e1_clean}, slots v14={len(v14['swap_vector'])}, v28={len(v28['swap_vector'])}")
results["E1"] = {"pass": e1_pass, "clean": e1_clean}

# E2: Difficulty-invariance index 1.00 for every generated variant
print("\n--- E2: Difficulty invariance check ---")
v42 = generate_centre_variant(exam_id, 42, questions, answer_key)
diff_indices = [v14["difficulty_index"], v28["difficulty_index"], v42["difficulty_index"]]
print(f"Difficulty indices: {diff_indices}")
e2_pass = all(d == 1.00 for d in diff_indices)
results["E2"] = {"pass": e2_pass, "indices": diff_indices}

# E3: POST leaked variant plaintext (centre 28) -> implicated 28, separation >= +50%
print("\n--- E3: Leak centre 28 plaintext ---")
text_28 = "\n".join(v28["variant_questions"])

r_insp = requests.post(
    f"{BASE_URL}/api/batch/inspect-variant",
    params={
        "leaked_text": text_28,
        "exam_id": exam_id,
        "centre_ids_raw": "14, 28, 42",
        "answer_key_text": answer_key,
    }
)
insp_data = r_insp.json() if r_insp.status_code == 200 else {}
print(f"Inspector result: {insp_data}")
imp = insp_data.get("implicated_centre_id")
sep = insp_data.get("separation_margin", 0.0)
e3_pass = (r_insp.status_code == 200) and (imp == 28) and (sep >= 50.0)
print(f"Implicated: {imp} (exp 28), Separation: {sep}% (exp >= 50%)")
results["E3"] = {"pass": e3_pass, "data": insp_data}

# E4: Same text with 2% char typos -> still 28
print("\n--- E4: 2% char typos ---")
chars = list(text_28)
num_typos = max(1, int(len(chars) * 0.02))
random.seed(42)
typo_indices = random.sample(range(len(chars)), num_typos)
for idx in typo_indices:
    if chars[idx].isalnum():
        chars[idx] = "x"
typo_text = "".join(chars)

r_insp_e4 = requests.post(
    f"{BASE_URL}/api/batch/inspect-variant",
    params={
        "leaked_text": typo_text,
        "exam_id": exam_id,
        "centre_ids_raw": "14, 28, 42",
        "answer_key_text": answer_key,
    }
)
insp_data_e4 = r_insp_e4.json() if r_insp_e4.status_code == 200 else {}
print(f"E4 result: {insp_data_e4}")
e4_pass = (r_insp_e4.status_code == 200) and (insp_data_e4.get("implicated_centre_id") == 28)
results["E4"] = {"pass": e4_pass, "data": insp_data_e4}

# E5: Master text (variants OFF) -> inspector does NOT mis-attribute a centre (UNKNOWN/low-conf)
print("\n--- E5: Master text (no variants) ---")
master_text = "\n".join(questions)
r_insp_e5 = requests.post(
    f"{BASE_URL}/api/batch/inspect-variant",
    params={
        "leaked_text": master_text,
        "exam_id": exam_id,
        "centre_ids_raw": "14, 28, 42",
        "answer_key_text": answer_key,
    }
)
insp_data_e5 = r_insp_e5.json() if r_insp_e5.status_code == 200 else {}
print(f"E5 result: {insp_data_e5}")
e5_pass = (insp_data_e5.get("status") == "INCONCLUSIVE") or (insp_data_e5.get("implicated_centre_id") is None)
print(f"E5 pass: {e5_pass}")
results["E5"] = {"pass": e5_pass, "data": insp_data_e5}

# E6: Numerical-only leak (numbers extracted from centre variant) -> still attributed
print("\n--- E6: Numerical-only leak ---")
numbers_only = " ".join(v28["number_map"].values())
print(f"Extracted numbers: {numbers_only}")
r_insp_e6 = requests.post(
    f"{BASE_URL}/api/batch/inspect-variant",
    params={
        "leaked_text": f"Values leaked: {numbers_only}",
        "exam_id": exam_id,
        "centre_ids_raw": "14, 28, 42",
        "answer_key_text": answer_key,
    }
)
insp_data_e6 = r_insp_e6.json() if r_insp_e6.status_code == 200 else {}
print(f"E6 result: {insp_data_e6}")
e6_pass = (r_insp_e6.status_code == 200) and (insp_data_e6.get("implicated_centre_id") == 28)
results["E6"] = {"pass": e6_pass, "data": insp_data_e6}

print("\n=== SUMMARY OF SECTION E ===")
for k, v in results.items():
    print(f"{k}: {'PASS' if v['pass'] else 'FAIL'} | {v}")
